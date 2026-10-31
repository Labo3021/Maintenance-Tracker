"""Import legacy inventories from CSV or Excel into assets.

Pipeline: read file -> find header row -> map columns (aliases, then Gemma for the
leftovers) -> parse each row into an AssetCreate -> report skipped rows/warnings.
"""

import csv
import io
import re
import unicodedata
from datetime import date, datetime, timedelta
from typing import Any, Optional

from pydantic import ValidationError

from ai.column_mapper import map_unknown_headers

from .schemas import AssetCreate, ColumnMapping, ImportIssue, ImportResult, MaintenanceSchedule, Part, TaskCreate
from .store import MatchIndex, store

PREVIEW_ROWS = 20


class ImportFileError(ValueError):
    pass


# --- Column mapping ----------------------------------------------------------

# Fields a spreadsheet column can fill, with descriptions (also shown to Gemma).
IMPORT_FIELDS: dict[str, str] = {
    "name": "Equipment or item name / description",
    "code": "Inventory code, SKU, catalog number, part number or asset tag",
    "category": "Equipment type, e.g. centrifuge, autoclave, hplc, freezer",
    "manufacturer": "Maker / brand / vendor",
    "model": "Model name or number",
    "serial_number": "Serial number or S/N",
    "location": "Room, lab, building, shelf, cabinet or bench",
    "status": "Condition / working state",
    "quantity_expected": "How many units there should be (theoretical / expected / par level)",
    "quantity_actual": "How many units were actually counted / on hand",
    "groups": "Kit, set or group the item belongs to (can be a comma-separated list)",
    "notes": "Free-text notes, comments or observations",
    "last_service_date": "Date of last maintenance / PM / service / repair visit",
    "service_interval_days": "How often it gets maintenance / PM (a number or text like '6 months')",
    "next_service_due": "Date the next service is due",
    "last_calibration_date": "Date of last calibration / certification / recertification / verification",
    "calibration_interval_days": "How often it is calibrated, certified or recertified",
    "next_calibration_due": "Date the next calibration / certification is due",
    "parts": "Consumable parts list",
}


def _fold(text: str) -> str:
    """Lowercase and strip accents: 'CÓDIGO' -> 'codigo', 'Ubicación' -> 'ubicacion'."""
    return "".join(c for c in unicodedata.normalize("NFKD", text.lower()) if not unicodedata.combining(c))


# Normalized header text -> field (English and Spanish). Normalization folds accents and drops punctuation.
ALIASES: dict[str, str] = {}
for _field, _names in {
    "name": ["name", "equipment", "equipment name", "instrument", "instrument name", "item", "item name",
             "asset", "asset name", "description", "equipment description", "device",
             "nombre", "descripcion", "equipo", "nombre del equipo", "instrumento", "articulo", "elemento",
             "material", "nombre del articulo"],
    "code": ["code", "item code", "sku", "part number", "part no", "part #", "catalog number", "catalog no",
             "cat no", "asset tag", "tag", "asset number", "inventory number", "inventory code",
             "codigo", "cod", "referencia", "ref", "numero de inventario", "codigo de inventario", "no inventario"],
    "category": ["category", "type", "equipment type", "asset type", "class", "instrument type",
                 "categoria", "tipo", "clase", "tipo de equipo"],
    "manufacturer": ["manufacturer", "mfr", "mfg", "make", "brand", "vendor", "maker", "supplier",
                     "fabricante", "marca", "proveedor"],
    "model": ["model", "model number", "model no", "model #", "model name", "modelo"],
    "serial_number": ["serial", "serial number", "serial no", "serial #", "s n", "sn", "serial num",
                      "equipment serial", "serie", "numero de serie", "n serie", "no serie", "nro serie",
                      "num serie", "n de serie"],
    "location": ["location", "room", "lab", "building", "site", "bench", "room number", "lab room", "shelf",
                 "shelves", "cabinet", "storage", "ubicacion", "estante", "estantes", "laboratorio", "sala",
                 "lugar", "aula", "bodega", "armario", "gabinete"],
    "status": ["status", "condition", "state", "working status", "operational status",
               "estado", "condicion", "estado actual"],
    "quantity_expected": ["expected", "expected quantity", "expected qty", "qty expected", "par level",
                          "target quantity", "theoretical quantity", "quantity expected",
                          "cantidad total supuesta", "cantidad supuesta", "cantidad esperada", "cantidad teorica",
                          "stock teorico", "cantidad inicial"],
    "quantity_actual": ["quantity", "qty", "count", "counted", "actual quantity", "actual qty", "qty on hand",
                        "on hand", "stock", "units", "quantity on hand", "actual count",
                        "cantidad", "cantidad real", "cantidad total real", "existencias", "unidades",
                        "stock real", "cantidad contada", "conteo"],
    "groups": ["kit", "kits", "kit name", "set", "group", "groups", "collection",
               "nombre del kit", "grupo", "conjunto", "coleccion"],
    "notes": ["notes", "note", "comments", "comment", "remarks", "observations",
              "nota", "notas", "comentario", "comentarios", "observacion", "observaciones"],
    "last_service_date": ["last service", "last service date", "last serviced", "last maintenance",
                          "last maintenance date", "last pm", "last pm date", "serviced on", "service date",
                          "date serviced", "maintenance date", "ultimo mantenimiento",
                          "fecha ultimo mantenimiento", "fecha de mantenimiento", "fecha mantenimiento"],
    "service_interval_days": ["service interval", "maintenance interval", "pm interval", "service frequency",
                              "maintenance frequency", "pm frequency", "service interval days",
                              "service interval months", "frequency", "interval",
                              "frecuencia", "frecuencia de mantenimiento", "periodicidad",
                              "periodicidad de mantenimiento"],
    "next_service_due": ["next service", "next service date", "next service due", "service due",
                         "next maintenance", "next pm", "next pm date", "pm due", "due date", "next due",
                         "proximo mantenimiento", "fecha proximo mantenimiento"],
    "last_calibration_date": ["last calibration", "last calibration date", "last calibrated", "calibration date",
                              "calibrated on", "last cal", "cal date", "last certification", "certification date",
                              "ultima calibracion", "fecha de calibracion", "fecha ultima calibracion"],
    "calibration_interval_days": ["calibration interval", "calibration frequency", "cal interval", "cal frequency",
                                  "calibration interval days", "calibration interval months",
                                  "recertification interval", "certification interval", "recert cycle",
                                  "frecuencia de calibracion", "periodicidad de calibracion"],
    "next_calibration_due": ["next calibration", "next calibration date", "calibration due", "cal due",
                             "next cal", "next cal date", "certification due", "proxima calibracion",
                             "fecha proxima calibracion"],
    "parts": ["parts", "consumables", "spare parts", "parts list", "repuestos", "consumibles"],
}.items():
    for _n in _names:
        ALIASES[re.sub(r"[^a-z0-9]+", " ", _fold(_n)).strip()] = _field


def _normalize_header(h: str) -> str:
    h = re.sub(r"\((s|es)\)", "", _fold(h))  # plural markers: "ESTANTE(S)" -> "estante"
    h = re.sub(r"\(.*?\)", lambda m: " " + m.group(0)[1:-1] + " ", h)  # "Interval (months)" -> "interval months"
    return re.sub(r"[^a-z0-9#]+", " ", h).replace("#", " #").strip()


def _alias(header: str) -> Optional[str]:
    norm = _normalize_header(header)
    if norm in ALIASES:
        return ALIASES[norm]
    simple = re.sub(r"[^a-z0-9]+", " ", norm).strip()
    if simple in ALIASES:
        return ALIASES[simple]
    # Drop unit words: "service interval days" / "pm interval months" -> "service interval"
    stripped = re.sub(r"\b(days?|months?|years?|weeks?|yrs?|mos?|dias?|meses|mes|anos?|semanas?)\b", "", simple)
    return ALIASES.get(re.sub(r"\s+", " ", stripped).strip())


def map_columns(headers: list[str], samples: list[list[Any]], use_ai: bool) -> tuple[list[ColumnMapping], str]:
    mappings: list[ColumnMapping] = []
    used: set[str] = set()
    for h in headers:
        field = _alias(h) if h else None
        if field and field not in used:
            used.add(field)
            mappings.append(ColumnMapping(header=h, field=field, method="alias"))
        else:
            mappings.append(ColumnMapping(header=h, field=None, method="unmapped"))

    unknown = [i for i, m in enumerate(mappings) if m.field is None and m.header]
    if not unknown:
        return mappings, "not_needed"
    if not use_ai:
        return mappings, "disabled"

    free_fields = {f: d for f, d in IMPORT_FIELDS.items() if f not in used}
    columns = {mappings[i].header: [str(r[i]) for r in samples if i < len(r) and r[i] not in (None, "")][:3]
               for i in unknown}
    suggestion, source = map_unknown_headers(columns, free_fields)
    for i in unknown:
        field = suggestion.get(mappings[i].header)
        if field in free_fields and field not in used:
            used.add(field)
            mappings[i] = ColumnMapping(header=mappings[i].header, field=field, method="ai")
    return mappings, source


# --- Reading files -----------------------------------------------------------


def _read_csv(data: bytes) -> list[list[Any]]:
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    return [row for row in csv.reader(io.StringIO(text), dialect)]


def _read_xlsx(data: bytes, sheet: Optional[str]) -> tuple[list[list[Any]], str]:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:
        raise ImportFileError(f"Could not open Excel file: {e}") from e
    if sheet and sheet not in wb.sheetnames:
        raise ImportFileError(f"Sheet '{sheet}' not found. Sheets: {', '.join(wb.sheetnames)}")
    ws = wb[sheet] if sheet else wb.worksheets[0]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()
    return rows, ws.title


def _is_blank(v: Any) -> bool:
    return v is None or (isinstance(v, str) and not v.strip())


# --- Value parsing -----------------------------------------------------------

DATE_FORMATS = {
    "mdy": ["%m/%d/%Y", "%m/%d/%y", "%m-%d-%Y", "%m-%d-%y", "%m.%d.%Y"],
    "dmy": ["%d/%m/%Y", "%d/%m/%y", "%d-%m-%Y", "%d-%m-%y", "%d.%m.%Y"],
}
UNAMBIGUOUS_FORMATS = ["%Y-%m-%d", "%Y/%m/%d", "%Y-%m-%d %H:%M:%S", "%d-%b-%Y", "%d %b %Y", "%b %d %Y",
                       "%b %d, %Y", "%B %d, %Y", "%B %d %Y", "%d %B %Y", "%b-%y", "%B %Y", "%b %Y"]


def parse_date(value: Any, date_order: str = "mdy") -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and 20000 < value < 80000:  # Excel serial date
        return date(1899, 12, 30) + timedelta(days=int(value))
    text = str(value).strip()
    if re.fullmatch(r"\d{5}(\.0+)?", text):
        return parse_date(float(text))
    for fmt in UNAMBIGUOUS_FORMATS + DATE_FORMATS[date_order]:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"unrecognized date '{text}'")


INTERVAL_WORDS = {
    "daily": 1, "weekly": 7, "biweekly": 14, "monthly": 30, "bimonthly": 60, "quarterly": 91,
    "semiannual": 182, "semiannually": 182, "biannual": 182, "biannually": 182, "twice a year": 182,
    "annual": 365, "annually": 365, "yearly": 365, "biennial": 730, "every other year": 730,
    "diario": 1, "semanal": 7, "quincenal": 15, "mensual": 30, "bimestral": 60, "trimestral": 91,
    "semestral": 182, "anual": 365, "bienal": 730,
}
UNIT_DAYS = {"d": 1, "day": 1, "w": 7, "wk": 7, "week": 7, "m": 30, "mo": 30, "mon": 30, "month": 30,
             "y": 365, "yr": 365, "year": 365, "dia": 1, "semana": 7, "mes": 30, "mese": 30, "ano": 365}


def _header_unit(header: str) -> int:
    h = _fold(header)
    if re.search(r"\b(years?|yrs?|anos?)\b", h):
        return 365
    if re.search(r"\b(months?|mos?|meses|mes)\b", h):
        return 30
    if re.search(r"\b(weeks?|semanas?)\b", h):
        return 7
    return 1


def parse_interval(value: Any, header: str) -> int:
    if isinstance(value, (int, float)):
        days = round(value * _header_unit(header))
    else:
        text = re.sub(r"[-_]", "", _fold(str(value).strip()))
        text = re.sub(r"^(every|cada)\s+", "", text)
        if text in INTERVAL_WORDS:
            return INTERVAL_WORDS[text]
        m = re.fullmatch(r"(\d+(?:\.\d+)?)\s*([a-z]*?)s?", text)
        if not m:
            raise ValueError(f"unrecognized interval '{value}'")
        n, unit = float(m.group(1)), m.group(2)
        if unit and unit not in UNIT_DAYS:
            raise ValueError(f"unrecognized interval unit in '{value}'")
        days = round(n * (UNIT_DAYS[unit] if unit else _header_unit(header)))
    if days < 1:
        raise ValueError(f"interval must be positive, got '{value}'")
    return days


STATUS_WORDS = {
    "operational": ["operational", "ok", "okay", "good", "working", "active", "in service", "in use", "available",
                    "functional", "running", "yes", "pass", "passed"],
    "needs_attention": ["needs attention", "attention", "needs repair", "needs service", "due", "overdue",
                        "warning", "degraded", "fair", "poor", "issue", "issues", "check", "maintenance"],
    "out_of_service": ["out of service", "oos", "broken", "down", "not working", "failed", "fail", "inactive",
                       "retired", "decommissioned", "offline", "red tagged", "do not use", "no"],
}
for _status, _words in {
    "operational": ["bueno", "buen estado", "operativo", "funciona", "funcional", "en uso", "activo", "disponible",
                    "si", "ok funciona"],
    "needs_attention": ["regular", "requiere mantenimiento", "en revision", "revisar", "danado parcialmente",
                        "mantenimiento", "pendiente"],
    "out_of_service": ["malo", "danado", "roto", "no funciona", "fuera de servicio", "en reparacion",
                       "dado de baja", "inactivo", "averiado"],
}.items():
    STATUS_WORDS[_status] += _words


def parse_status(value: Any) -> str:
    text = re.sub(r"[^a-z ]+", " ", _fold(str(value))).strip()
    text = re.sub(r"\s+", " ", text)
    for status, words in STATUS_WORDS.items():
        if text in words:
            return status
    for status in ("out_of_service", "needs_attention", "operational"):  # most severe first
        if any(re.search(rf"\b{re.escape(w)}\b", text) for w in STATUS_WORDS[status] if len(w) > 2):
            return status
    raise ValueError(f"unrecognized status '{value}'")


def parse_parts(value: Any, date_order: str) -> list[Part]:
    """'Name | part# | installed date | lifespan' entries separated by ';' (only name is required)."""
    parts = []
    for entry in str(value).split(";"):
        bits = [b.strip() for b in entry.split("|")]
        if not bits[0]:
            continue
        part = Part(name=bits[0])
        if len(bits) > 1 and bits[1]:
            part.part_number = bits[1]
        if len(bits) > 2 and bits[2]:
            part.installed_date = parse_date(bits[2], date_order)
        if len(bits) > 3 and bits[3]:
            part.lifespan_days = parse_interval(bits[3], "days")
        parts.append(part)
    return parts


def parse_quantity(value: Any) -> int:
    if isinstance(value, (int, float)):
        if value < 0 or (isinstance(value, float) and not value.is_integer()):
            raise ValueError(f"not a whole, non-negative quantity: '{value}'")
        return int(value)
    m = re.fullmatch(r"\s*(\d+)(?:[.,]0+)?\s*[a-zA-Z.]*\s*", str(value))  # "12", "12.0", "12 u", "12 pcs"
    if not m:
        raise ValueError(f"unrecognized quantity '{value}'")
    return int(m.group(1))


PLACEHOLDERS = {"0", "-", "--", "n/a", "na", "none", "ninguno", "ninguna", "no aplica", "sin kit", "s/n", "?"}


def parse_groups(value: Any) -> list[str]:
    """Comma/semicolon-separated groups. Not '/', which often appears inside names ("EST / COMPLETO")."""
    seen: list[str] = []
    for g in re.split(r"[,;]", _text(value)):
        g = re.sub(r"\s+", " ", g).strip()
        if g and _fold(g) not in PLACEHOLDERS and g.lower() not in (x.lower() for x in seen):
            seen.append(g)
    return seen


def _text(value: Any) -> str:
    if isinstance(value, float) and value.is_integer():
        value = int(value)  # Excel turns serials/model numbers like 1290 into 1290.0
    return str(value).strip()


# --- Row -> asset ------------------------------------------------------------


def _parse_row(
    row: list[Any], columns: list[ColumnMapping], date_order: str, row_no: int, warnings: list[ImportIssue]
) -> tuple[AssetCreate, list[TaskCreate]]:
    values: dict[str, Any] = {}
    headers: dict[str, str] = {}
    for i, col in enumerate(columns):
        if col.field and i < len(row) and not _is_blank(row[i]):
            values[col.field] = row[i]
            headers[col.field] = col.header

    out: dict[str, Any] = {}
    for field in ("name", "code", "category", "manufacturer", "model", "serial_number", "location", "notes"):
        if field in values:
            out[field] = _text(values[field])
    # A bare shelf number like "1,2" means nothing on its own; keep the column title with it.
    if "location" in out and re.fullmatch(r"[\d\s,;/-]+", out["location"]):
        label = re.sub(r"\(.*?\)", "", headers["location"]).strip().capitalize()
        out["location"] = f"{label} {out['location']}"
    if "category" in out:
        out["category"] = re.sub(r"[^a-z0-9]+", "_", out["category"].lower()).strip("_") or "general"

    if not out.get("name"):
        fallback = " ".join(filter(None, [out.get("manufacturer"), out.get("model")])) or out.get("code", "")
        if not fallback:
            raise ValueError("no name, manufacturer/model or code in this row")
        out["name"] = fallback
        warnings.append(ImportIssue(row=row_no, message=f"No name; used '{fallback}'"))

    def soft(field: str, parser, *args):
        """Parse an optional field; a bad value becomes a warning, not a skipped row."""
        if field not in values:
            return None
        try:
            return parser(values[field], *args)
        except ValueError as e:
            warnings.append(ImportIssue(row=row_no, message=f"{headers[field]}: {e} (ignored)"))
            return None

    if (status := soft("status", parse_status)) is not None:
        out["status"] = status
    if (parts := soft("parts", parse_parts, date_order)) is not None:
        out["parts"] = parts
    for field in ("quantity_expected", "quantity_actual"):
        if (qty := soft(field, parse_quantity)) is not None:
            out[field] = qty
    if (groups := soft("groups", parse_groups)) is not None:
        out["groups"] = groups

    tasks: list[TaskCreate] = []
    for kind in ("service", "calibration"):
        last_f, interval_f, next_f = (f"last_{kind}_date", f"{kind}_interval_days", f"next_{kind}_due")
        last = soft(last_f, parse_date, date_order)
        interval = soft(interval_f, parse_interval, headers.get(interval_f, ""))
        nxt = soft(next_f, parse_date, date_order)
        if last is None and nxt is not None and interval is not None:
            last = nxt - timedelta(days=interval)  # keep the schedule the sheet implies
        elif nxt is not None and interval is None:
            # A due date with no interval can't become a recurring schedule; keep it as a task.
            tasks.append(TaskCreate(title=f"Scheduled {kind} (imported)", due_date=nxt, priority="medium"))
        if last is not None:
            out[last_f] = last
        if interval is not None:
            out[interval_f] = interval

    try:
        return AssetCreate(**out), tasks
    except ValidationError as e:
        raise ValueError("; ".join(f"{'.'.join(map(str, err['loc']))}: {err['msg']}" for err in e.errors())) from e


# --- Entry point -------------------------------------------------------------


def import_inventory(
    filename: str,
    data: bytes,
    *,
    mode: str = "append",
    dry_run: bool = False,
    date_order: str = "mdy",
    use_ai: bool = True,
    sheet: Optional[str] = None,
    stock_count_days: int = 180,
) -> ImportResult:
    name = filename.lower()
    if name.endswith(".csv") or name.endswith(".txt"):
        file_type, rows, sheet_name = "csv", _read_csv(data), None
    elif name.endswith(".xlsx") or name.endswith(".xlsm"):
        file_type = "xlsx"
        rows, sheet_name = _read_xlsx(data, sheet)
    elif name.endswith(".xls"):
        raise ImportFileError("Old .xls files aren't supported. In Excel use File > Save As > .xlsx or .csv")
    else:
        raise ImportFileError("Upload a .csv or .xlsx file")

    # Header = first row with 2+ filled cells (skips title rows like "Lab Inventory 2019").
    header_idx = next((i for i, r in enumerate(rows) if sum(not _is_blank(v) for v in r) >= 2), None)
    if header_idx is None:
        raise ImportFileError("Couldn't find a header row (need at least 2 column titles)")
    headers = [_text(h) if not _is_blank(h) else "" for h in rows[header_idx]]
    body = [(header_idx + 2 + i, r) for i, r in enumerate(rows[header_idx + 1 :]) if not all(_is_blank(v) for v in r)]

    columns, ai_source = map_columns(headers, [r for _, r in body[:5]], use_ai)
    if not any(c.field in ("name", "model", "manufacturer", "code") for c in columns):
        raise ImportFileError(
            "No column looks like an item name, code, model or manufacturer. "
            f"Headers found: {', '.join(h for h in headers if h)}"
        )

    parsed: list[tuple[AssetCreate, list[TaskCreate]]] = []
    skipped: list[ImportIssue] = []
    warnings: list[ImportIssue] = []
    for row_no, row in body:
        try:
            parsed.append(_parse_row(row, columns, date_order, row_no, warnings))
        except ValueError as e:
            skipped.append(ImportIssue(row=row_no, message=str(e)))

    kits = _kit_assets([a for a, _ in parsed], stock_count_days)
    rows_to_save = parsed + [(k, []) for k in kits]

    if dry_run:
        index = MatchIndex([] if mode == "replace" else store.list_assets())
        updated = 0
        for i, (asset, _) in enumerate(rows_to_save):
            if index.find(asset):
                updated += 1
            else:
                index.add(asset, f"new-{i}")  # later duplicate rows in the same file update this one
        created = len(rows_to_save) - updated
    else:
        created, updated = store.import_assets(rows_to_save, mode)

    return ImportResult(
        dry_run=dry_run,
        mode=mode,
        file_type=file_type,
        sheet=sheet_name,
        header_row=header_idx + 1,
        columns=columns,
        ai_mapping=ai_source,
        total_rows=len(body),
        created=created,
        updated=updated,
        skipped=skipped,
        warnings=warnings,
        preview=[a for a, _ in parsed[:PREVIEW_ROWS]],
        stock_mismatches=sum(
            1 for a, _ in parsed
            if a.quantity_expected is not None and a.quantity_actual is not None
            and a.quantity_actual < a.quantity_expected
        ),
        kits_created=len(kits),
    )


def _kit_assets(items: list[AssetCreate], stock_count_days: int) -> list[AssetCreate]:
    """One asset per kit/group with a recurring physical stock count. The import itself counts
    as the latest count when the sheet has counted quantities; otherwise the first count is due now."""
    if stock_count_days <= 0:
        return []
    members: dict[str, list[AssetCreate]] = {}
    for item in items:
        for g in item.groups:
            members.setdefault(g, []).append(item)
    counted = any(i.quantity_actual is not None for i in items)
    kits = []
    for name, group in members.items():
        missing = sum(
            max(i.quantity_expected - i.quantity_actual, 0)
            for i in group if i.quantity_expected is not None and i.quantity_actual is not None
        )
        kits.append(AssetCreate(
            name=name,
            code="kit:" + re.sub(r"[^a-z0-9]+", "-", _fold(name)).strip("-"),
            category="kit",
            notes=f"{len(group)} item types" + (f"; {missing} units missing at import" if missing else ""),
            schedules=[MaintenanceSchedule(
                title="Physical stock count", kind="inspect", interval_days=stock_count_days,
                last_done=date.today() if counted else None,
                notes="Count every item in the kit and update counted quantities.",
            )],
        ))
    return kits
